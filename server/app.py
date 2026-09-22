"""Helios: a single-user, loopback-only interview workbench."""
from __future__ import annotations

import asyncio
import base64
import contextlib
import hashlib
import io
import json
import logging
import math
import os
from pathlib import Path
import secrets
import shutil
import sqlite3
import tempfile
import time
from typing import Literal
from urllib.parse import urlparse
import uuid
import wave
import zipfile

import httpx
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, ConfigDict, Field
from starlette.staticfiles import StaticFiles

from .store import Store, ROOT, DATA_DIR, COLLECTIONS, DEFAULT_SETTINGS, now
from . import ai, speech, codex_provider

log = logging.getLogger("helios")
store = Store()
COOKIE = "helios_session"
TOKEN = secrets.token_urlsafe(32)
MAX_UPLOAD = 1024 * 1024 * 1024
task_tokens: dict[str, str] = {}
live_sessions: set[str] = set()
active_tasks: set[str] = set()
readiness_cache: tuple[float, dict] = (0, {})


def private_settings():
    settings = store.settings()
    try:
        import keyring
        settings["api_key"] = keyring.get_password("Helios", "provider-api-key") or ""
    except Exception:
        settings["api_key"] = ""
    return settings


def require(collection, record_id):
    try:
        return store.get(collection, record_id)
    except KeyError as exc:
        raise HTTPException(404, str(exc).strip("'")) from exc


def public_task(task):
    return {k: v for k, v in task.items() if k != "settings_snapshot"}


def create_task(kind, target_id, **fields):
    existing = next((t for t in store.list("tasks", kind=kind, target_id=target_id)
                     if t["status"] in ("queued", "running")), None)
    if existing:
        return public_task(existing)
    task = store.put("tasks", {"kind": kind, "target_id": target_id, "status": "queued",
                              "progress": 0, "message": "Waiting to start", "error": None,
                              "settings_snapshot": store.settings(), **fields})
    return public_task(task)


def check_cancel(task_id):
    if require("tasks", task_id)["status"] == "cancelled":
        raise RuntimeError("Processing cancelled")


def progress(task_id, value, message):
    check_cancel(task_id)
    store.update("tasks", task_id, {"progress": max(0, min(1, float(value))), "message": message})


def session_segments(session_id):
    return sorted(store.list("segments", session_id=session_id), key=lambda s: (s["start"], s["created_at"]))


def invalidate_reports(session_id):
    for report in store.list("assessments", session_id=session_id):
        store.update("assessments", report["id"], {"stale": True})


async def execute_task(task):
    tid, kind, target = task["id"], task["kind"], task["target_id"]
    settings = task.get("settings_snapshot", store.settings())
    settings["api_key"] = private_settings().get("api_key", "")
    # Current offline preference always overrides older queued cloud jobs.
    if store.settings()["offline"]:
        settings["offline"] = True
    progress(tid, .05, "Starting")
    if kind in ("requirements", "preparation"):
        job = require("jobs", target)
        progress(tid, .15, "Reading the job and your experience")
        if kind == "requirements":
            result = await ai.requirements(job, settings)
            check_cancel(tid)
            store.update("jobs", target, {"requirements": result})
        else:
            result = await ai.prepare(job, settings)
            check_cancel(tid)
            store.put("preparations", {"job_id": target, "plan": result}, tid)
    elif kind == "assessment":
        session = require("sessions", target)
        segments = session_segments(target)
        candidates = [s for s in segments if s["role"] == "candidate" and s["text"].strip()]
        if not candidates:
            raise ValueError("Identify your answers in the transcript before requesting feedback.")
        if session["source"] == "upload" and not session.get("speakers_confirmed"):
            raise ValueError("Confirm which speaker is you before requesting feedback.")
        progress(tid, .2, "Reviewing your answers against the job requirements")
        result = await ai.assess(require("jobs", session["job_id"]), session, segments, settings)
        check_cancel(tid)
        current_revision = require("sessions", target).get("transcript_revision", 0)
        store.put("assessments", {"session_id": target, "report": result, "stale": current_revision != session.get("transcript_revision", 0),
                                 "provider": settings["provider"], "transcript_revision": session.get("transcript_revision", 0)}, tid)
        for idx, drill in enumerate(result.get("drills", [])):
            store.put("drills", {**drill, "job_id": session["job_id"], "session_id": target, "completed": False}, f"{tid}-{idx}")
        store.update("sessions", target, {"status": "reviewed"})
    elif kind == "import":
        session = require("sessions", target)
        path = media_dir(target) / session["original_filename"]
        def callback(value, message):
            progress(tid, value, message)
        segments = await asyncio.to_thread(speech.process_upload, path, settings, callback)
        check_cancel(tid)
        for old in store.list("segments", session_id=target):
            store.delete("segments", old["id"])
        for idx, segment in enumerate(segments):
            store.put("segments", {**segment, "session_id": target}, f"{tid}-{idx}")
        store.update("sessions", target, {"status": "needs-speakers", "speakers_confirmed": False,
                     "duration_seconds": max((s.get("end", 0) for s in segments), default=0)})
    elif kind == "model-download":
        component = task["component"]
        if component == "llm":
            await pull_ollama(task)
        else:
            token = task_tokens.pop(tid, "")
            paths = await asyncio.to_thread(speech.download_model, component, token, settings,
                store.directory / "models", lambda value, message: progress(tid, value, message))
            check_cancel(tid)
            store.save_settings({"speech": paths})
    else:
        raise ValueError("Unknown processing task")
    check_cancel(tid)
    store.update("tasks", tid, {"status": "completed", "progress": 1, "message": "Complete", "error": None})
    global readiness_cache
    readiness_cache = (0, {})


async def pull_ollama(task):
    settings = store.settings()
    provider = settings["provider"]
    if provider["kind"] != "ollama":
        raise ValueError("Select Ollama in Settings to download its model, or install a model in your selected runtime.")
    parsed = urlparse(provider["base_url"])
    if parsed.hostname not in ("localhost", "127.0.0.1", "::1"):
        raise ValueError("Model downloads can only target a local Ollama server.")
    origin = f"{parsed.scheme}://{parsed.netloc}"
    model = provider.get("model") or "qwen3:8b"
    if "cloud" in model.lower():
        raise ValueError("Choose a local model such as qwen3:8b.")
    async with httpx.AsyncClient(timeout=httpx.Timeout(1800, connect=5), trust_env=False, follow_redirects=False) as client:
        async with client.stream("POST", origin + "/api/pull", json={"name": model, "stream": True}) as response:
            response.raise_for_status()
            async for line in response.aiter_lines():
                if not line:
                    continue
                event = json.loads(line)
                if event.get("error"):
                    raise ValueError(event["error"])
                total = event.get("total", 0)
                progress(task["id"], event.get("completed", 0) / total if total else .1, event.get("status", "Downloading"))


async def worker():
    for task in store.list("tasks"):
        if task["status"] == "running":
            store.update("tasks", task["id"], {"status": "queued", "message": "Resuming after restart"})
    while True:
        queued = list(reversed(store.list("tasks", status="queued")))
        # Active interviews take priority over CPU-heavy imported recordings.
        task = next((t for t in queued if not (live_sessions and t["kind"] in ("import", "model-download"))), None)
        if task:
            active_tasks.add(task["id"])
            store.update("tasks", task["id"], {"status": "running"})
            try:
                await execute_task(task)
            except asyncio.CancelledError:
                raise
            except Exception as exc:
                current = require("tasks", task["id"])
                if current["status"] != "cancelled":
                    message = str(exc)[:1200] or "Processing failed. Please retry."
                    # Never put a download credential in an error message.
                    if task.get("component") == "diarization" and "token" in message.lower():
                        message = "Speaker model access failed. Accept the model conditions and provide a valid download token in Settings."
                    store.update("tasks", task["id"], {"status": "failed", "error": message, "message": message})
            finally:
                active_tasks.discard(task["id"])
        else:
            await asyncio.sleep(.5)


@contextlib.asynccontextmanager
async def lifespan(app):
    background = asyncio.create_task(worker())
    yield
    background.cancel()
    with contextlib.suppress(asyncio.CancelledError):
        await background


app = FastAPI(title="Helios", lifespan=lifespan, docs_url=None, redoc_url=None)


def local_host(host):
    name = host.split(":")[0] if not host.startswith("[") else host.split("]")[0] + "]"
    return name.lower() in ("localhost", "127.0.0.1", "[::1]")


def valid_origin(origin):
    return origin in ("http://localhost:8765", "http://127.0.0.1:8765", "http://localhost:5173", "http://127.0.0.1:5173")


@app.middleware("http")
async def local_security(request: Request, call_next):
    if not local_host(request.headers.get("host", "")):
        return JSONResponse({"detail": "Helios is available only on localhost."}, status_code=403)
    origin = request.headers.get("origin")
    if origin and not valid_origin(origin):
        return JSONResponse({"detail": "Cross-origin requests are not allowed."}, status_code=403)
    if request.headers.get("sec-fetch-site") == "cross-site":
        return JSONResponse({"detail": "Open Helios directly on localhost."}, status_code=403)
    if request.url.path.startswith("/api/") and request.url.path not in ("/api/bootstrap", "/api/health"):
        if not secrets.compare_digest(request.cookies.get(COOKIE, ""), TOKEN):
            return JSONResponse({"detail": "Reload Helios to reconnect to the local service."}, status_code=401)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Cache-Control"] = "no-store" if request.url.path.startswith("/api") else "no-cache"
    return response


@app.exception_handler(ValueError)
async def value_error(request, exc):
    return JSONResponse({"detail": str(exc)}, status_code=400)


@app.get("/api/health")
def health():
    return {"ok": True, "name": "Helios", "version": "0.1.0", "pid": os.getpid(),
            "instance": hashlib.sha256(str(ROOT).lower().encode()).hexdigest()[:16]}


async def get_readiness():
    global readiness_cache
    if time.monotonic() - readiness_cache[0] < 15:
        return readiness_cache[1]
    settings = private_settings()
    result = await asyncio.to_thread(speech.readiness, settings)
    try:
        status = await asyncio.wait_for(ai.test_provider(settings), timeout=4)
        models = status.get("models", [])
        selected = settings["provider"].get("model", "")
        ready = status.get("ok", False) and bool(selected)
        if models and selected not in models:
            ready = False
        result["llm"] = {"ready": ready, "message": status.get("message", "Model connected") if ready else status.get("message", "Select an installed model in Settings.")}
    except Exception as exc:
        result["llm"] = {"ready": False, "message": str(exc) or "The language model is not connected. Open Settings to connect Ollama or another provider."}
    vad = ROOT / "dist" / "vad"
    result["vad"] = {"ready": any(vad.glob("*.onnx")) and any(vad.glob("*.wasm")),
                     "message": "Local voice detection assets available" if any(vad.glob("*.onnx")) else "Build the frontend to install local voice detection assets."}
    readiness_cache = (time.monotonic(), result)
    return result


@app.get("/api/bootstrap")
async def bootstrap():
    data = {name: store.list(name) for name in ("jobs", "rounds", "sessions", "preparations", "assessments", "drills")}
    data.update(tasks=[public_task(t) for t in store.list("tasks")], settings=store.settings(), readiness=await get_readiness())
    response = JSONResponse(data)
    response.set_cookie(COOKIE, TOKEN, httponly=True, samesite="strict", max_age=86400 * 30)
    return response


@app.get("/api/readiness")
async def readiness():
    return await get_readiness()


@app.get("/api/settings")
def get_settings():
    return store.settings()


class SettingsPatch(BaseModel):
    model_config = ConfigDict(extra="forbid")
    offline: bool | None = None
    provider: dict | None = None
    pause_ms: int | None = Field(None, ge=400, le=10000)
    ignore_mic_while_speaking: bool | None = None
    speech: dict | None = None
    api_key: str | None = None


@app.patch("/api/settings")
def save_settings(body: SettingsPatch):
    global readiness_cache
    patch = body.model_dump(exclude_none=True)
    key = patch.pop("api_key", None)
    if key is not None:
        try:
            import keyring
            if key:
                keyring.set_password("Helios", "provider-api-key", key)
            else:
                with contextlib.suppress(keyring.errors.PasswordDeleteError):
                    keyring.delete_password("Helios", "provider-api-key")
        except Exception as exc:
            raise HTTPException(503, "The operating system credential store is unavailable; the API key was not saved.") from exc
    if "provider" in patch:
        patch["provider"] = {k: v for k, v in patch["provider"].items() if k in {"kind", "base_url", "model"}}
    if "speech" in patch:
        patch["speech"] = {k: v for k, v in patch["speech"].items() if k in {"stt_path", "tts_path", "diarization_path"}}
    readiness_cache = (0, {})
    return store.save_settings(patch)


@app.post("/api/providers/test")
async def test_provider(body: dict):
    settings = private_settings()
    settings["provider"] = body.get("provider", {k: v for k, v in body.items() if k in {"kind", "base_url", "model"}})
    if "offline" in body:
        settings["offline"] = bool(body["offline"])
    if "api_key" in body:
        settings["api_key"] = body["api_key"]
    try:
        return await ai.test_provider(settings)
    except Exception as exc:
        return {"ok": False, "models": [], "message": str(exc)}


@app.get("/api/codex/status")
async def codex_status():
    return await codex_provider.status(private_settings())


@app.post("/api/codex/login")
async def codex_login():
    return await codex_provider.login(private_settings())


@app.post("/api/codex/logout")
async def codex_logout():
    return await codex_provider.logout(private_settings())


class JobInput(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    company: str = Field("", max_length=200)
    description: str = Field("", max_length=100000)
    resume_text: str = Field("", max_length=100000)
    requirements: list[str] = Field(default_factory=list)
    seniority: str = "Mid-level"
    status: str = "preparing"


@app.get("/api/jobs")
def jobs():
    return store.list("jobs")


@app.post("/api/jobs")
def add_job(body: JobInput):
    return store.put("jobs", body.model_dump())


@app.patch("/api/jobs/{job_id}")
def edit_job(job_id: str, body: dict):
    previous = require("jobs", job_id)
    return store.update("jobs", job_id, JobInput(**{**previous, **body}).model_dump())


@app.delete("/api/jobs/{job_id}")
def delete_job(job_id: str):
    require("jobs", job_id)
    for s in store.list("sessions", job_id=job_id):
        ensure_idle_session(s["id"])
    if any(t["target_id"] == job_id and (t["status"] in ("queued", "running") or t["id"] in active_tasks) for t in store.list("tasks")):
        raise HTTPException(409, "Cancel the job's processing task before deleting it.")
    for s in store.list("sessions", job_id=job_id):
        delete_session(s["id"])
    for collection in ("rounds", "preparations", "drills"):
        for record in store.list(collection, job_id=job_id):
            store.delete(collection, record["id"])
    store.delete("jobs", job_id)
    return {"ok": True}


@app.post("/api/jobs/{job_id}/extract")
def extract_requirements(job_id: str):
    require("jobs", job_id)
    return create_task("requirements", job_id)


@app.post("/api/jobs/{job_id}/prepare")
def prepare_job(job_id: str):
    require("jobs", job_id)
    return create_task("preparation", job_id)


@app.post("/api/documents/extract")
async def extract_document(file: UploadFile = File(...)):
    data = await file.read(15 * 1024 * 1024 + 1)
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(413, "Documents must be smaller than 15 MB.")
    suffix = Path(file.filename or "").suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:100])
        elif suffix == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        elif suffix in (".txt", ".md"):
            text = data.decode("utf-8-sig")
        else:
            raise ValueError("Use a PDF, DOCX, or UTF-8 text file.")
    except Exception as exc:
        raise HTTPException(400, "Could not read this document. Use a text-based PDF, DOCX, or paste the text directly.") from exc
    if not text.strip():
        raise HTTPException(400, "This document has no extractable text. Paste its text instead; scanned PDFs need OCR first.")
    return {"text": text[:100000]}


class RoundInput(BaseModel):
    job_id: str
    title: str = Field(min_length=1, max_length=200)
    kind: str = "screening"
    scheduled_at: str = ""
    status: str = "upcoming"
    outcome: str = ""
    notes: str = ""


@app.get("/api/rounds")
def rounds():
    return store.list("rounds")


@app.post("/api/rounds")
def add_round(body: RoundInput):
    require("jobs", body.job_id)
    return store.put("rounds", body.model_dump())


@app.patch("/api/rounds/{round_id}")
def edit_round(round_id: str, body: dict):
    validated = RoundInput(**{**require("rounds", round_id), **body})
    require("jobs", validated.job_id)
    return store.update("rounds", round_id, validated.model_dump())


@app.delete("/api/rounds/{round_id}")
def delete_round(round_id: str):
    require("rounds", round_id)
    store.delete("rounds", round_id)
    for session in store.list("sessions", round_id=round_id):
        store.update("sessions", session["id"], {"round_id": None})
    return {"ok": True}


class SessionInput(BaseModel):
    job_id: str
    round_id: str | None = None
    title: str = "Practice interview"
    kind: str = "role-specific"
    mode: Literal["learning", "mock"] = "mock"
    difficulty: str = "adaptive"
    duration_minutes: int = Field(30, ge=5, le=120)


@app.get("/api/sessions")
def sessions():
    return store.list("sessions")


@app.post("/api/sessions")
def add_session(body: SessionInput):
    require("jobs", body.job_id)
    if body.round_id and require("rounds", body.round_id)["job_id"] != body.job_id:
        raise ValueError("The selected round belongs to a different job.")
    return store.put("sessions", {**body.model_dump(), "status": "ready", "source": "practice",
        "duration_seconds": 0, "provider": store.settings()["provider"], "next_seq": 0, "transcript_revision": 0})


@app.get("/api/sessions/{session_id}")
def session_detail(session_id: str):
    return {"session": require("sessions", session_id), "segments": session_segments(session_id),
            "assessments": store.list("assessments", session_id=session_id)}


def media_dir(session_id):
    # Never accept a user-controlled filesystem path as a record identifier.
    try:
        uuid.UUID(session_id)
    except ValueError:
        raise HTTPException(400, "Invalid session identifier")
    directory = store.directory / "media" / session_id
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def ensure_idle_session(session_id):
    if session_id in live_sessions or any(t["target_id"] == session_id and (t["status"] in ("queued", "running") or t["id"] in active_tasks) for t in store.list("tasks")):
        raise HTTPException(409, "Stop the interview and cancel active processing before deleting this recording.")


@app.delete("/api/sessions/{session_id}")
def delete_session(session_id: str):
    require("sessions", session_id)
    ensure_idle_session(session_id)
    for collection in ("segments", "assessments", "drills"):
        for record in store.list(collection, session_id=session_id):
            store.delete(collection, record["id"])
    directory = media_dir(session_id)
    if directory.resolve().parent == (store.directory / "media").resolve():
        shutil.rmtree(directory)
    store.delete("sessions", session_id)
    return {"ok": True}


@app.post("/api/sessions/{session_id}/end")
def end_session(session_id: str):
    require("sessions", session_id)
    store.update("sessions", session_id, {"status": "completed"})
    if not any(s["role"] == "candidate" for s in session_segments(session_id)):
        return {"status": "completed", "message": "Session saved. Add answers before requesting feedback."}
    return create_task("assessment", session_id)


@app.post("/api/sessions/{session_id}/assess")
def assess_session(session_id: str):
    if require("sessions", session_id)["status"] == "active":
        raise HTTPException(409, "End the interview before requesting feedback.")
    return create_task("assessment", session_id)


@app.post("/api/sessions/{session_id}/hint")
async def hint_session(session_id: str):
    session = require("sessions", session_id)
    if session["mode"] != "learning":
        raise HTTPException(400, "Hints are available in learning sessions. Mock interview feedback is provided after the round.")
    try:
        return {"text": await ai.hint(require("jobs", session["job_id"]), session, session_segments(session_id), private_settings())}
    except Exception as exc:
        raise HTTPException(503, str(exc)) from exc


@app.patch("/api/segments/{segment_id}")
def edit_segment(segment_id: str, body: dict):
    segment = require("segments", segment_id)
    fields = {k: v for k, v in body.items() if k in ("text", "role", "start", "end")}
    merged = {**segment, **fields}
    if merged["role"] not in ("candidate", "interviewer", "unknown", "other"):
        raise ValueError("Select a valid speaker role.")
    if not isinstance(merged["text"], str) or len(merged["text"]) > 50000:
        raise ValueError("Transcript text is too long.")
    start, end = float(merged["start"]), float(merged["end"])
    if not math.isfinite(start) or not math.isfinite(end) or not 0 <= start <= end:
        raise ValueError("Transcript timestamps must be positive and in order.")
    fields.update(start=start, end=end)
    record = store.update("segments", segment_id, fields)
    session = require("sessions", segment["session_id"])
    store.update("sessions", session["id"], {"transcript_revision": session.get("transcript_revision", 0) + 1})
    invalidate_reports(session["id"])
    return record


@app.post("/api/sessions/{session_id}/speakers")
def map_speakers(session_id: str, body: dict):
    previous = require("sessions", session_id)
    mapping = body.get("mapping", {})
    if any(role not in ("candidate", "interviewer", "unknown", "other") for role in mapping.values()):
        raise ValueError("Select a valid role for each speaker.")
    with store.lock:
        segments = session_segments(session_id)
        if not any(mapping.get(s.get("speaker"), s["role"]) == "candidate" for s in segments):
            raise ValueError("Identify at least one of your answers before confirming speakers.")
        for segment in segments:
            if segment.get("speaker") in mapping:
                store.update("segments", segment["id"], {"role": mapping[segment["speaker"]]})
        session = store.update("sessions", session_id, {"speakers_confirmed": True, "status": "completed",
            "transcript_revision": previous.get("transcript_revision", 0) + 1})
        invalidate_reports(session_id)
    return session


@app.patch("/api/drills/{drill_id}")
def edit_drill(drill_id: str, body: dict):
    require("drills", drill_id)
    if not isinstance(body.get("completed"), bool):
        raise ValueError("A drill must be marked complete or incomplete.")
    return store.update("drills", drill_id, {"completed": body["completed"]})


@app.post("/api/tasks/{task_id}/cancel")
def cancel_task(task_id: str):
    task = require("tasks", task_id)
    if task["status"] in ("queued", "running"):
        task_tokens.pop(task_id, None)
        task = store.update("tasks", task_id, {"status": "cancelled", "message": "Cancelled; saved recordings are preserved."})
    return public_task(task)


@app.post("/api/tasks/{task_id}/retry")
def retry_task(task_id: str):
    task = require("tasks", task_id)
    if task_id in active_tasks:
        raise HTTPException(409, "The previous attempt is still stopping. Retry in a moment.")
    if task["status"] not in ("failed", "cancelled"):
        raise HTTPException(409, "Only failed or cancelled tasks can be retried.")
    if task.get("component") == "diarization":
        raise HTTPException(400, "Start the speaker model download again from Settings to supply a new temporary access token.")
    return public_task(store.update("tasks", task_id, {"status": "queued", "progress": 0, "error": None,
        "message": "Waiting to retry", "settings_snapshot": store.settings()}))


@app.post("/api/models/download")
def download_model(body: dict):
    component = body.get("component")
    if component not in ("stt", "tts", "diarization", "llm"):
        raise ValueError("Unknown model component")
    task = create_task("model-download", component, component=component)
    if body.get("token"):
        task_tokens[task["id"]] = body["token"]
    return task


@app.post("/api/sessions/{session_id}/chunks")
async def save_chunk(session_id: str, request: Request, seq: int, duration_seconds: float = 0):
    session = require("sessions", session_id)
    if session["source"] != "practice" or not 0 <= seq < 100000:
        raise ValueError("Invalid recording chunk")
    data = await request.body()
    if len(data) > 8 * 1024 * 1024 or not data:
        raise HTTPException(413, "Recording chunks must be between 1 byte and 8 MB.")
    try:
        with wave.open(io.BytesIO(data), "rb") as wav:
            if wav.getsampwidth() != 2 or wav.getnchannels() not in (1, 2) or not 8000 <= wav.getframerate() <= 96000:
                raise ValueError("Unsupported recording format")
    except (wave.Error, EOFError) as exc:
        raise HTTPException(400, "Recording chunks must be valid PCM WAV audio.") from exc
    folder = media_dir(session_id) / "chunks"
    folder.mkdir(exist_ok=True)
    target = folder / f"{seq:08d}.wav"
    if target.exists():
        if hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(data).digest():
            raise HTTPException(409, "A different audio chunk already uses this sequence number.")
    else:
        temp = target.with_suffix(".part")
        temp.write_bytes(data)
        temp.replace(target)
    store.update("sessions", session_id, {"next_seq": max(session.get("next_seq", 0), seq + 1),
                 "duration_seconds": max(session.get("duration_seconds", 0), min(duration_seconds, 12 * 3600))})
    return {"ok": True, "next_seq": seq + 1}


def assemble_recording(session_id):
    directory = media_dir(session_id)
    chunks = sorted((directory / "chunks").glob("*.wav"))
    if not chunks:
        return None
    numbers = [int(p.stem) for p in chunks]
    if numbers != list(range(max(numbers) + 1)):
        raise HTTPException(409, "Some audio chunks are still waiting to upload. Keep Helios open and retry saving.")
    target = directory / "recording.wav"
    temporary = directory / "recording.tmp.wav"
    # Decode each independently to support resumed sessions with different device rates.
    import numpy as np
    with wave.open(str(temporary), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(16000)
        for chunk in chunks:
            audio = speech._decode_audio(chunk)
            pcm = (np.clip(audio, -1, 1) * 32767).astype("<i2")
            output.writeframes(pcm.tobytes())
    temporary.replace(target)
    with wave.open(str(target), "rb") as audio:
        duration = audio.getnframes() / audio.getframerate()
    return {"recording_filename": target.name, "recording_url": f"/api/sessions/{session_id}/recording", "duration_seconds": duration}


@app.post("/api/sessions/{session_id}/finalize")
async def finalize_recording(session_id: str, body: dict):
    session = require("sessions", session_id)
    fields = await asyncio.to_thread(assemble_recording, session_id)
    return store.update("sessions", session_id, fields) if fields else session


@app.get("/api/sessions/{session_id}/recording")
def recording(session_id: str):
    session = require("sessions", session_id)
    filename = session.get("recording_filename") or session.get("original_filename")
    if not filename or Path(filename).name != filename:
        raise HTTPException(404, "This session has no saved recording yet.")
    path = media_dir(session_id) / filename
    if not path.is_file():
        raise HTTPException(404, "The recording file is missing.")
    return FileResponse(path, filename=f"Helios-{session_id[:8]}{path.suffix}", content_disposition_type="inline")


@app.post("/api/imports")
async def import_recording(file: UploadFile = File(...), job_id: str = Form(...), title: str = Form("Imported interview")):
    require("jobs", job_id)
    extension = Path(file.filename or "").suffix.lower()
    if extension not in (".wav", ".mp3", ".m4a", ".mp4", ".webm", ".ogg", ".flac", ".aac", ".mov", ".mkv"):
        raise ValueError("Choose an audio or video recording: WAV, MP3, M4A, MP4, WebM, OGG, FLAC, AAC, MOV, or MKV.")
    session = store.put("sessions", {"job_id": job_id, "title": title[:200], "kind": "role-specific",
        "mode": "mock", "difficulty": "adaptive", "duration_minutes": 30, "duration_seconds": 0,
        "status": "processing", "source": "upload", "speakers_confirmed": False, "transcript_revision": 0,
        "original_filename": "original" + extension, "next_seq": 0})
    path = media_dir(session["id"]) / session["original_filename"]
    length = 0
    try:
        with path.open("wb") as output:
            while chunk := await file.read(1024 * 1024):
                length += len(chunk)
                if length > MAX_UPLOAD:
                    raise HTTPException(413, "Recordings must be smaller than 1 GB.")
                output.write(chunk)
        if not length:
            raise ValueError("This recording is empty.")
    except Exception:
        path.unlink(missing_ok=True)
        store.delete("sessions", session["id"])
        raise
    session = store.update("sessions", session["id"], {"recording_url": f"/api/sessions/{session['id']}/recording"})
    return {"session": session, "task": create_task("import", session["id"])}


@app.websocket("/api/sessions/{session_id}/live")
async def live(websocket: WebSocket, session_id: str):
    if not valid_origin(websocket.headers.get("origin", "")) or not local_host(websocket.headers.get("host", "")) or not secrets.compare_digest(websocket.cookies.get(COOKIE, ""), TOKEN):
        await websocket.close(code=1008)
        return
    try:
        session = store.get("sessions", session_id)
    except KeyError:
        await websocket.close(code=1008)
        return
    if session["source"] != "practice" or session_id in live_sessions:
        await websocket.close(code=1008, reason="This session is already open, or is an imported recording.")
        return
    await websocket.accept()
    live_sessions.add(session_id)
    store.update("sessions", session_id, {"status": "active"})
    generation = None
    current_request = ""
    settings = private_settings()
    settings["provider"] = session.get("provider", settings["provider"])
    job = require("jobs", session["job_id"])

    async def emit(payload, rid):
        if rid == current_request:
            await websocket.send_json({**payload, "request_id": rid})

    async def run_turn(message, rid):
        try:
            if store.settings()["offline"]:
                settings["offline"] = True
            # Request IDs make answer retries idempotent even after a reconnect.
            candidate_id = str(uuid.uuid5(uuid.UUID(session_id), rid + ":candidate"))
            if message["type"] in ("answer", "audio"):
                text = str(message.get("text", "")).strip()
                if message["type"] == "audio":
                    await emit({"type": "state", "value": "transcribing"}, rid)
                    encoded = message.get("audio", "")
                    if len(encoded) > 28 * 1024 * 1024:
                        raise ValueError("This answer is too long. Use Send Answer more often or type the answer.")
                    audio = base64.b64decode(encoded, validate=True)
                    parts = await asyncio.to_thread(speech.transcribe, audio, settings)
                    text = " ".join(p["text"] for p in parts).strip()
                if not text:
                    raise ValueError("No speech was detected. Try again or type your answer.")
                start = max(0, float(message.get("start", 0)))
                end = max(start, float(message.get("end", start)))
                segment = store.put("segments", {"session_id": session_id, "role": "candidate", "speaker": "You",
                    "text": text[:50000], "start": start, "end": end}, candidate_id)
                await emit({"type": "segment", "segment": segment}, rid)
            elif message["type"] == "start":
                previous = session_segments(session_id)
                if previous and previous[-1]["role"] == "interviewer":
                    # A reconnect does not invent a new unanswered question.
                    await emit({"type": "question", "text": previous[-1]["text"], "audio": None, "segment": previous[-1]}, rid)
                    await emit({"type": "state", "value": "listening"}, rid)
                    return
            await emit({"type": "state", "value": "thinking"}, rid)
            segments = session_segments(session_id)
            current_session = require("sessions", session_id)
            current_session["duration_seconds"] = max(current_session.get("duration_seconds", 0), float(message.get("end", 0)))
            turn = await ai.next_turn(job, current_session, segments, settings)
            if rid != current_request:
                return
            question = turn["question"]
            wav = None
            warning = None
            try:
                wav = await asyncio.to_thread(speech.synthesize, question, settings)
            except Exception as exc:
                warning = str(exc)
            start = max(float(message.get("end", 0)), max((s["end"] for s in segments), default=0))
            duration = 0
            if wav:
                with wave.open(io.BytesIO(wav), "rb") as voice:
                    duration = voice.getnframes() / voice.getframerate()
            segment = store.put("segments", {"session_id": session_id, "role": "interviewer", "speaker": "Helios",
                       "text": question, "start": start, "end": start + duration}, str(uuid.uuid5(uuid.UUID(session_id), rid + ":question")))
            await emit({"type": "segment", "segment": segment}, rid)
            await emit({"type": "question", "text": question, "audio": base64.b64encode(wav).decode() if wav else None,
                        "segment": segment, "warning": warning}, rid)
            await emit({"type": "state", "value": "speaking" if wav else "listening"}, rid)
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            await emit({"type": "error", "message": str(exc) or "Helios could not complete this turn."}, rid)
            await emit({"type": "state", "value": "listening"}, rid)

    try:
        while True:
            message = await websocket.receive_json()
            kind = message.get("type")
            if kind == "playback":
                segment = require("segments", message.get("segment_id", ""))
                if segment["session_id"] == session_id and segment["role"] == "interviewer":
                    start = max(0, float(message.get("start", 0)))
                    store.update("segments", segment["id"], {"start": start, "end": max(start, float(message.get("end", start)))})
                continue
            if kind not in ("start", "answer", "audio", "interrupt", "pause", "end"):
                await websocket.send_json({"type": "error", "message": "Unknown session event"})
                continue
            if generation and not generation.done():
                generation.cancel()
                with contextlib.suppress(asyncio.CancelledError):
                    await generation
            current_request = str(message.get("request_id") or uuid.uuid4())[:200]
            if kind in ("interrupt", "pause", "end"):
                await websocket.send_json({"type": "cancelled", "request_id": current_request})
                continue
            generation = asyncio.create_task(run_turn(message, current_request))
    except (WebSocketDisconnect, RuntimeError):
        pass
    finally:
        if generation and not generation.done():
            generation.cancel()
            with contextlib.suppress(asyncio.CancelledError, RuntimeError):
                await generation
        live_sessions.discard(session_id)
        if require("sessions", session_id)["status"] == "active":
            store.update("sessions", session_id, {"status": "paused"})


@app.get("/api/backup")
def backup():
    if live_sessions or active_tasks:
        raise HTTPException(409, "Finish the active interview or processing task before creating a consistent backup.")
    tempdir = Path(tempfile.mkdtemp(prefix="backup-", dir=store.directory / "tmp"))
    database = tempdir / "helios.sqlite3"
    store.snapshot(database)
    archive = tempdir / "helios-backup.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as output:
        output.writestr("manifest.json", json.dumps({"format": "helios-backup", "version": 1, "created_at": now()}))
        output.write(database, "helios.sqlite3")
        for path in (store.directory / "media").rglob("*"):
            if path.is_file() and not path.is_symlink() and path.suffix != ".part":
                output.write(path, str(path.relative_to(store.directory)).replace("\\", "/"))
    from starlette.background import BackgroundTask
    return FileResponse(archive, filename="helios-backup.zip", background=BackgroundTask(shutil.rmtree, tempdir))


@app.post("/api/restore")
async def restore(file: UploadFile = File(...)):
    if live_sessions or active_tasks or any(t["status"] in ("running", "queued") for t in store.list("tasks")):
        raise HTTPException(409, "Finish or cancel active interviews and processing before restoring a backup.")
    stage = Path(tempfile.mkdtemp(prefix="restore-", dir=store.directory / "tmp"))
    try:
        archive_path = stage / "upload.zip"
        length = 0
        with archive_path.open("wb") as output:
            while chunk := await file.read(1024 * 1024):
                length += len(chunk)
                if length > 2 * 1024**3:
                    raise HTTPException(413, "Backups must be smaller than 2 GB.")
                output.write(chunk)
        extracted = stage / "contents"
        extracted.mkdir()
        with zipfile.ZipFile(archive_path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 4 * 1024**3:
                raise ValueError("Expanded backup exceeds 4 GB.")
            names = set()
            for info in archive.infolist():
                name = info.filename.replace("\\", "/")
                target = (extracted / name).resolve()
                if not target.is_relative_to(extracted.resolve()) or ":" in name or ".." in Path(name).parts or ((info.external_attr >> 16) & 0o170000) == 0o120000:
                    raise ValueError("Backup contains an unsafe file path.")
                if name in names:
                    raise ValueError("Backup contains duplicate files.")
                names.add(name)
                if name not in ("manifest.json", "helios.sqlite3") and not name.startswith("media/"):
                    raise ValueError("This archive contains files outside the Helios backup format.")
                if info.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                else:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    with archive.open(info) as source, target.open("wb") as output:
                        shutil.copyfileobj(source, output)
        manifest = json.loads((extracted / "manifest.json").read_text())
        if manifest.get("format") != "helios-backup" or manifest.get("version") != 1:
            raise ValueError("This backup format is not supported.")
        database = extracted / "helios.sqlite3"
        with contextlib.closing(sqlite3.connect(database)) as con, con:
            if con.execute("PRAGMA quick_check").fetchone()[0] != "ok" or con.execute("PRAGMA user_version").fetchone()[0] != 1:
                raise ValueError("The backup database is invalid.")
            objects = con.execute("SELECT type,name FROM sqlite_master WHERE name NOT LIKE 'sqlite_%'").fetchall()
            if set(objects) != {("table", "records"), ("table", "preferences")}:
                raise ValueError("The backup contains an unsupported database schema.")
            # Validate logical contents before replacing any current data.
            for collection, record_id, raw in con.execute("SELECT collection,id,data FROM records"):
                record = json.loads(raw)
                if collection not in COLLECTIONS or not isinstance(record, dict) or record.get("id") != record_id or not record.get("created_at"):
                    raise ValueError("The backup contains invalid records.")
                if collection in ("jobs", "rounds", "sessions"):
                    uuid.UUID(record_id)
                if collection == "jobs":
                    JobInput(**record)
                elif collection == "rounds":
                    RoundInput(**record)
                elif collection == "sessions":
                    SessionInput(**record)
                elif collection == "segments":
                    if not isinstance(record.get("text"), str) or record.get("role") not in ("candidate", "interviewer", "unknown", "other"):
                        raise ValueError("The backup contains an invalid transcript.")
                    start, end = float(record["start"]), float(record["end"])
                    if not math.isfinite(start) or not math.isfinite(end) or not 0 <= start <= end:
                        raise ValueError("The backup contains invalid transcript timestamps.")
                elif collection == "preparations":
                    from .ai_schemas import PreparationPlan
                    PreparationPlan.model_validate(record["plan"])
                elif collection == "assessments":
                    from .ai_schemas import AssessmentReport
                    AssessmentReport.model_validate(record["report"])
            settings = json.loads(con.execute("SELECT data FROM preferences WHERE id=1").fetchone()[0])
            if not isinstance(settings, dict) or not isinstance(settings.get("provider"), dict) or not isinstance(settings.get("speech"), dict):
                raise ValueError("The backup contains invalid settings.")
            settings = {key: settings.get(key, default) for key, default in DEFAULT_SETTINGS.items()}
            SettingsPatch(**settings)
            if any(not isinstance(settings["provider"].get(key), str) for key in ("kind", "base_url", "model")):
                raise ValueError("The backup contains an invalid model connection.")
            settings["provider"] = {key: settings["provider"][key] for key in ("kind", "base_url", "model")}
            settings["speech"] = {key: str(settings["speech"].get(key, "")) for key in ("stt_path", "tts_path", "diarization_path")}
            settings["offline"] = True
            con.execute("UPDATE preferences SET data=? WHERE id=1", (json.dumps(settings),))
            rows = con.execute("SELECT id,data FROM records WHERE collection='tasks'").fetchall()
            for rid, raw in rows:
                task = json.loads(raw)
                if task["status"] in ("running", "queued"):
                    task.update(status="cancelled", message="Restored task; retry when ready.")
                    con.execute("UPDATE records SET data=? WHERE collection='tasks' AND id=?", (json.dumps(task), rid))
        # Keep the previous state for recovery; SQLite backup avoids WAL replacement hazards.
        previous = store.directory / "tmp" / f"before-restore-{uuid.uuid4().hex}"
        previous.mkdir()
        store.snapshot(previous / "helios.sqlite3")
        shutil.copytree(store.directory / "media", previous / "media")
        new_media = extracted / "media"
        new_media.mkdir(exist_ok=True)
        with store.lock:
            with contextlib.closing(sqlite3.connect(database)) as source, store.connect() as destination:
                source.backup(destination)
            shutil.rmtree(store.directory / "media")
            shutil.copytree(new_media, store.directory / "media")
        global readiness_cache
        readiness_cache = (0, {})
        return {"ok": True, "message": "Backup restored. Offline mode is enabled; reconnect your chosen models in Settings."}
    except (zipfile.BadZipFile, KeyError, TypeError, sqlite3.Error, OSError, json.JSONDecodeError) as exc:
        raise HTTPException(400, "This is not a readable Helios backup.") from exc
    finally:
        shutil.rmtree(stage, ignore_errors=True)


# Registered after API routes so local assets can never shadow the API.
if (ROOT / "dist" / "index.html").exists():
    app.mount("/", StaticFiles(directory=ROOT / "dist", html=True), name="frontend")
else:
    @app.get("/")
    def not_built():
        return JSONResponse({"detail": "Build Helios with npm run build, then restart the launcher."}, status_code=503)
